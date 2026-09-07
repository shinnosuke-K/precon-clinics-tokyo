#!/usr/bin/env python3
"""登録医療機関一覧のExcelからindex.htmlを更新するスクリプト。

使い方:
    python3 scripts/build.py [xlsxファイル] [htmlファイル]

引数を省略すると data/ 内で名前順最後のExcelと、リポジトリ直下の index.html を使う。
Excelの構成（1枚目: 登録医療機関リスト、2枚目以降: 医療機関ごとの詳細シート）を読み、
index.html 内の `const DATA = [...];` の行と「令和N年N月N日時点」の日付3箇所を置き換える。
ページのHTML・CSS・JS本体には手を加えない。

依存: openpyxl (pip install openpyxl)
"""
import datetime
import json
import re
import sys
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
CIRCLE = ("○", "〇")

KB_MAP = {
    "男性及び女性の検査、助言・相談が可能": "both",
    "男性の検査、助言・相談が可能": "m",
    "女性の検査、助言・相談が可能": "f",
}

KEY_ORDER = ["n", "nm", "note", "ct", "ad", "tel", "dt", "as", "kb", "kb2",
             "sp", "rv", "cs", "fee", "must", "tm", "tf"]


def norm(s):
    """全角スペースを半角に直し、前後の空白を落とす。"""
    return s.replace("　", " ").strip() if isinstance(s, str) else s


def amount(v):
    """金額セルの値を文字列化する。数値はそのまま、文字列（「（上記金額に含む）」等）は正規化のみ。"""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return norm(str(v))


def serial_to_iso(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    return (datetime.date(1899, 12, 30) + datetime.timedelta(days=int(v))).isoformat()


def is_circle(v):
    return isinstance(v, str) and v.strip() in CIRCLE


def parse_tests(rows, start):
    """男性検査/女性検査セクションを読む。

    検査名がある行の金額セルをその検査の金額とし（なければ空文字）、
    サブ項目（選択項目）は○が付いているものだけ採用する。
    金額もサブ項目の○も無い検査は、その医療機関では実施なしとして除外する。
    返り値: ([[検査名, 金額, [[サブ項目, 金額], ...]], ...], 次に読む行番号)
    """
    tests = []
    i = start
    while i < len(rows):
        r = rows[i]
        b, c, flag, e, h = r[1], r[2], r[3], r[4], r[7]
        if i > start and (b is not None or all(x is None for x in (c, flag, e, h))):
            break
        if c is not None:
            tests.append([norm(c), amount(h), [], is_circle(flag) or h is not None])
        if e is not None and tests and is_circle(flag):
            tests[-1][2].append([norm(e), amount(h)])
            tests[-1][3] = True
        i += 1
    return [[name, amt if amt is not None else "", subs]
            for name, amt, subs, included in tests if included], i


def parse_detail(ws):
    rows = list(ws.iter_rows(values_only=True))
    d = {"as": "", "tel": "", "rv": [], "cs": [], "kb2": "",
         "fee": ["", "", ""], "must": [], "tm": [], "tf": []}
    for v in rows[0]:
        if isinstance(v, str) and "時点" in v:
            d["as"] = v.replace("時点", "").strip()
    section = None
    i = 0
    while i < len(rows):
        r = rows[i]
        b = r[1]
        if b == "電話番号":
            d["tel"] = norm(r[3]) or ""
        elif b == "予約方法":
            section = "rv"
        elif b == "助言・相談の実施方法":
            section = "cs"
        elif b == "登録区分":
            section = "kb2"
        elif b in ("初診料", "再診料", "助言・相談料"):
            d["fee"][("初診料", "再診料", "助言・相談料").index(b)] = amount(r[7]) or ""
            section = None
        elif b == "必須検査":
            section = "must"
        elif b == "男性検査":
            d["tm"], i = parse_tests(rows, i)
            continue
        elif b == "女性検査":
            d["tf"], i = parse_tests(rows, i)
            continue
        if section in ("rv", "cs") and r[4] is not None:
            if is_circle(r[3]):
                d[section].append(norm(r[4]))
        elif section == "kb2" and r[4] in KB_MAP:
            if is_circle(r[3]):
                d["kb2"] = KB_MAP[r[4]]
        elif section == "must" and r[2] is not None:
            d["must"].append([norm(r[2]), amount(r[7]) or ""])
        elif section == "must" and all(x is None for x in r[1:8]):
            section = None
        i += 1
    return d


def build(xlsx_path):
    """Excelを読み、(レコードのリスト, 「令和N年N月N日時点」の文字列) を返す。"""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    list_ws = wb[wb.sheetnames[0]]
    detail_sheets = wb.sheetnames[1:]
    all_rows = list(list_ws.iter_rows(values_only=True))

    asof = None
    for v in all_rows[0]:
        if isinstance(v, str) and (m := re.search(r"令和\d+年\d+月\d+日時点", v)):
            asof = m.group(0)
    if not asof:
        sys.exit("一覧シートの見出しから「令和N年N月N日時点」を取得できませんでした")

    rows = [r for r in all_rows if isinstance(r[0], int)]
    if len(rows) != len(detail_sheets):
        sys.exit(f"一覧の行数 {len(rows)} と詳細シート数 {len(detail_sheets)} が一致しません")

    records = []
    for row, sheet in zip(rows, detail_sheets):
        n, dt, ct, nm_cell, ad, kb, sp = row[:7]
        nm, _, note = norm(nm_cell).partition("\n")
        rec = {"n": n, "nm": nm.strip(), "note": note.strip(), "ct": ct,
               "ad": norm(ad), "dt": serial_to_iso(dt),
               "kb": KB_MAP[kb.strip()], "sp": is_circle(sp)}
        rec.update(parse_detail(wb[sheet]))
        records.append({k: rec[k] for k in KEY_ORDER})
    return records, asof


def main():
    if len(sys.argv) > 1:
        xlsx = Path(sys.argv[1])
    else:
        candidates = sorted((REPO / "data").glob("*.xlsx"))
        if not candidates:
            sys.exit("data/ にExcelファイルがありません")
        xlsx = candidates[-1]
    html_path = Path(sys.argv[2]) if len(sys.argv) > 2 else REPO / "index.html"

    records, asof = build(xlsx)
    data_js = json.dumps(records, ensure_ascii=False, separators=(",", ":"))

    html = html_path.read_text(encoding="utf-8")
    html, n_data = re.subn(r"^const DATA = \[.*\];$", lambda m: f"const DATA = {data_js};",
                           html, flags=re.M)
    if n_data != 1:
        sys.exit(f"index.html 内の const DATA 行が {n_data} 箇所でした（1箇所のはず）")
    html, n_date = re.subn(r"令和\d+年\d+月\d+日時点", asof, html)

    html_path.write_text(html, encoding="utf-8")
    print(f"{xlsx.name} から {len(records)} 機関を読み込み、{html_path.name} を更新しました"
          f"（{asof}、日付の置換 {n_date} 箇所）")


if __name__ == "__main__":
    main()
